#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
md_to_docx.py —— Markdown → WPS/Office 友好 .docx 阅读版生成器（派生视图，非事实来源）

用途
    「Markdown 宽表在 WPS 里超页宽、右侧列被裁掉」（`ISSUE-020` 起因）的处置手段之一：
    把主件 .md 导出为**横向 A4** .docx —— 表格按列数自适应字号与列宽，不再横向溢出。

纪律（本工程约定，勿违）
    1. **主件仍是唯一事实来源**；本件为**派生视图**，不具独立受控效力，文件内已写明。
    2. **只做版式转换，不改一字内容**；唯一「新增」是自动派生并明确标注的《需求状态一览》。
    3. 产物落 `CM-002` §3 豁免清单第 7 类（机器生成件、不取号、不改名）。
    4. **必须可重跑**：主件升版后重跑本脚本即可刷新阅读版。

用法
    python md_to_docx.py [主件.md] [输出.docx]
    缺省：主件 00-需求/软件级需求/REQ-005_软件级需求文档.md
          产物 00-需求/软件级需求/REQ-005_软件级需求文档_阅读版.docx
"""
import io
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
DEFAULT_MD = os.path.join(ROOT, '00-需求', '软件级需求', 'REQ-005_软件级需求文档.md')
DEFAULT_OUT = os.path.join(ROOT, '00-需求', '软件级需求', 'REQ-005_软件级需求文档_阅读版.docx')

CJK = '微软雅黑'
MONO = 'Consolas'

HEAD_SIZE = {1: 16, 2: 13, 3: 11.5, 4: 10.5, 5: 10, 6: 10}


# ----------------------------------------------------------------- 通用小工具
def cell_font_size(n_col):
    """按列数自适应字号：列越多字越小（横向 A4 下 12 列仍可读）。"""
    if n_col <= 3:
        return 9.5
    if n_col <= 6:
        return 8.5
    if n_col <= 8:
        return 7.5
    return 6.5


def is_row(line):
    s = line.strip()
    return len(s) > 1 and s.startswith('|') and s.endswith('|')


def is_sep(line):
    return bool(re.match(r'^\|[\s:\-|]+\|$', line.strip()))


def split_row(line):
    body = line.strip().strip('|')
    parts = re.split(r'(?<!\\)\|', body)
    return [p.replace('\\|', '|').strip() for p in parts]


def set_run(run, size=None, bold=False, mono=False, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    f = run.font
    f.name = MONO if mono else CJK
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), MONO if mono else CJK)
    rf.set(qn('w:hAnsi'), MONO if mono else CJK)
    rf.set(qn('w:eastAsia'), CJK)


TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]*`|\{\{[^}]*\}\})')


def add_inline(par, text, size, base_bold=False):
    """渲染行内 **粗体**、`代码`、{{占位符}}（粗体内可嵌套代码段）。"""
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            for sub in re.split(r'(`[^`]*`)', tok[2:-2]):
                if not sub:
                    continue
                if sub.startswith('`') and sub.endswith('`') and len(sub) >= 2:
                    set_run(par.add_run(sub[1:-1]), size=size - 0.5, mono=True, bold=True)
                else:
                    set_run(par.add_run(sub), size=size, bold=True)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            set_run(par.add_run(tok[1:-1]), size=size - 0.5, mono=True)
        elif tok.startswith('{{'):
            set_run(par.add_run(tok), size=size, mono=True, color=RGBColor(0xB0, 0x30, 0x30))
        else:
            set_run(par.add_run(tok), size=size, bold=base_bold)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hex_color)
    tcpr.append(sh)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trpr.append(el)


def set_table_full_width(table, n_col):
    """表宽 = 页面可用宽度 100%，并设等宽列。"""
    tblpr = table._tbl.tblPr
    w = OxmlElement('w:tblW')
    w.set(qn('w:type'), 'pct')
    w.set(qn('w:w'), '5000')
    tblpr.append(w)
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblpr.append(lay)
    avail = 27.3  # 横向 A4 29.7cm - 左右各 1.2cm
    for row in table.rows:
        for c in row.cells:
            c.width = Cm(avail / n_col)


# ----------------------------------------------------------------- 主流程
def parse(md_text):
    return md_text.replace('\r\n', '\n').split('\n')


def collect_states(lines):
    """预扫 §2 功能需求：抽取「编号 / 需求名称 / 状态」派生视图数据。"""
    out = []
    in_sec2 = False
    fence = False
    i = 0
    while i < len(lines):
        l = lines[i]
        if re.match(r'^\s*```', l):
            fence = not fence
            i += 1
            continue
        if not fence and l.startswith('## '):
            in_sec2 = l.strip().startswith('## 2')
        if not fence and in_sec2 and is_row(l) and i + 1 < len(lines) and is_sep(lines[i + 1]):
            head = split_row(l)
            if '状态' in head and '编号' in head[0]:
                si = head.index('状态')
                j = i + 2
                while j < len(lines) and is_row(lines[j]):
                    c = split_row(lines[j])
                    if len(c) > si:
                        code = c[0].replace('**', '').strip()
                        if code and not code.startswith('编号'):
                            out.append((code, c[1].replace('**', '').strip(), c[si].replace('**', '').strip()))
                    j += 1
                i = j
                continue
        i += 1
    return out


def build(md_path, out_path):
    raw = io.open(md_path, encoding='utf-8', newline='').read()
    lines = parse(raw)
    states = collect_states(lines)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    for attr, v in (('left_margin', 1.2), ('right_margin', 1.2), ('top_margin', 1.2), ('bottom_margin', 1.2)):
        setattr(sec, attr, Cm(v))

    st = doc.styles['Normal']
    st.font.name = CJK
    st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)

    src_name = os.path.basename(md_path)
    out_name = os.path.basename(out_path)

    def banner():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        set_run(p.add_run('⚠️ 本件为《%s》（主件）的「阅读版派生视图」' % src_name), size=9.5, bold=True,
                color=RGBColor(0xC0, 0x30, 0x30))
        add_inline(p, '　由 03-工程代码/delivery/tools/md_to_docx.py 自**主件**自动生成、**可随时重跑**'
                      '（主件升版后重跑即刷新）。**不具独立受控效力**，一切以主件为准。'
                      '页面取 **横向 A4**，表格按列数自适应字号与列宽，以便在 WPS / Office 中阅读；正文未改动一字。', 9)

    def add_states():
        h = doc.add_heading('', level=2)
        set_run(h.add_run('阅读版附加视图 · 需求状态一览'), size=13, bold=True)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, '⚠️ 本表为**自动派生视图**（由 §2 各表的「编号／需求名称／状态」三列汇总），'
                      '用于一眼看清状态；**以主件 §2 各表为准**，不构成主件内容。', 8.5)
        n = 3
        t = doc.add_table(rows=1, cols=n)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0]
        for k, name in enumerate(('编号', '需求名称', '状态')):
            c = hdr.cells[k]
            c.text = ''
            add_inline(c.paragraphs[0], name, cell_font_size(n), base_bold=True)
            shade(c, 'D9E2F3')
        repeat_header(hdr)
        for code, nm, sta in states:
            r = t.add_row()
            for k, val in enumerate((code, nm, sta)):
                c = r.cells[k]
                c.text = ''
                add_inline(c.paragraphs[0], val, cell_font_size(n))
        set_table_full_width(t, n)
        doc.add_paragraph()

    fence = False
    inserted_states = False
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if re.match(r'^\s*```', line):
            fence = not fence
            i += 1
            continue
        if fence:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(line), size=8, mono=True)
            i += 1
            continue

        # 表格
        if is_row(line) and i + 1 < len(lines) and is_sep(lines[i + 1]):
            head = split_row(line)
            n = len(head)
            fs = cell_font_size(n)
            body = []
            j = i + 2
            while j < len(lines) and is_row(lines[j]):
                body.append(split_row(lines[j]))
                j += 1
            t = doc.add_table(rows=1, cols=n)
            t.style = 'Table Grid'
            hdr = t.rows[0]
            for k in range(n):
                c = hdr.cells[k]
                c.text = ''
                add_inline(c.paragraphs[0], head[k] if k < len(head) else '', fs, base_bold=True)
                shade(c, 'D9E2F3')
            repeat_header(hdr)
            for row in body:
                r = t.add_row()
                for k in range(n):
                    c = r.cells[k]
                    c.text = ''
                    add_inline(c.paragraphs[0], row[k] if k < len(row) else '', fs)
            set_table_full_width(t, n)
            doc.add_paragraph()
            i = j
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            lvl = len(m.group(1))
            if not inserted_states and lvl == 2 and stripped.startswith('## 1'):
                add_states()
                inserted_states = True
            h = doc.add_heading('', level=lvl)
            h.paragraph_format.space_before = Pt(10 if lvl <= 2 else 6)
            h.paragraph_format.space_after = Pt(4)
            set_run(h.add_run(re.sub(r'\*\*|`', '', m.group(2))), size=HEAD_SIZE.get(lvl, 10), bold=True)
            if lvl == 1:
                banner()
            i += 1
            continue

        # 分隔线
        if stripped in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            ppr = p._p.get_or_add_pPr()
            bd = OxmlElement('w:pBdr')
            bt = OxmlElement('w:bottom')
            bt.set(qn('w:val'), 'single')
            bt.set(qn('w:sz'), '6')
            bt.set(qn('w:color'), 'BFBFBF')
            bd.append(bt)
            ppr.append(bd)
            i += 1
            continue

        # 引用
        if stripped.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, re.sub(r'^>\s?', '', stripped), 8.5)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
            i += 1
            continue

        # 列表
        m = re.match(r'^(\s*)([-*+])\s+(.*)$', line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 + len(m.group(1)) * 0.25)
            p.paragraph_format.space_after = Pt(1)
            set_run(p.add_run('• '), size=9)
            add_inline(p, m.group(3), 9)
            i += 1
            continue
        m = re.match(r'^(\s*)(\d+)[.)]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 + len(m.group(1)) * 0.25)
            p.paragraph_format.space_after = Pt(1)
            set_run(p.add_run(m.group(2) + '. '), size=9)
            add_inline(p, m.group(3), 9)
            i += 1
            continue

        # 空行
        if stripped == '':
            i += 1
            continue

        # 正文
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, stripped, 9)
        i += 1

    # 页脚
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run('阅读版派生视图（横向 A4）· 主件：%s · 由 md_to_docx.py 自动生成' % src_name), size=7.5,
            color=RGBColor(0x80, 0x80, 0x80))

    doc.save(out_path)
    print('✅ 已生成 %s' % out_path)
    print('   主件 %d 行 / %d 字符；派生状态一览 %d 条' % (len(lines), len(raw), len(states)))


if __name__ == '__main__':
    a = sys.argv[1:]
    build(a[0] if a else DEFAULT_MD, a[1] if len(a) > 1 else DEFAULT_OUT)
